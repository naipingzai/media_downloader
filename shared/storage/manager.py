"""统一数据存储管理器 —— 支持 CSV / XLSX / SQLite / JSON。"""
from pathlib import Path
from typing import TYPE_CHECKING

from ..core.constants import VOLUME

if TYPE_CHECKING:
    pass

__all__ = ["StorageManager"]


class StorageManager:
    """管理采集数据的持久化存储。"""

    FORMATS = ("csv", "xlsx", "sql", "json", "")

    def __init__(self, storage_format: str = "", root: Path | None = None):
        self.format = storage_format if storage_format in self.FORMATS else ""
        self.root = root or VOLUME

    def get_output_dir(self, folder_name: str = "Data") -> Path:
        """获取数据存储目录，不存在则创建。"""
        path = self.root / folder_name
        path.mkdir(parents=True, exist_ok=True)
        return path

    async def save(self, data: list[dict], filename: str, **kwargs) -> str | None:
        """将数据保存到对应格式的文件。"""
        if not self.format or not data:
            return None
        output_dir = self.get_output_dir()
        match self.format:
            case "csv":
                return await self._save_csv(data, output_dir / f"{filename}.csv")
            case "xlsx":
                return self._save_xlsx(data, output_dir / f"{filename}.xlsx")
            case "sql":
                return await self._save_sql(data, filename, output_dir)
            case "json":
                return await self._save_json(data, output_dir / f"{filename}.json")
        return None

    async def _save_csv(self, data: list[dict], path: Path) -> str:
        import aiofiles
        import csv

        async with aiofiles.open(path, "w", encoding="utf-8-sig", newline="") as f:
            if data:
                writer = csv.DictWriter(f, fieldnames=data[0].keys())
                writer.writeheader()
                writer.writerows(data)
        return str(path)

    def _save_xlsx(self, data: list[dict], path: Path) -> str:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        if data:
            ws.append(list(data[0].keys()))
            for row in data:
                ws.append(list(row.values()))
        wb.save(path)
        return str(path)

    async def _save_sql(self, data: list[dict], table_name: str, output_dir: Path) -> str:
        import aiosqlite

        db_path = output_dir / "data.db"
        async with aiosqlite.connect(db_path) as db:
            if data:
                columns = ", ".join(data[0].keys())
                placeholders = ", ".join(["?"] * len(data[0]))
                await db.execute(
                    f"CREATE TABLE IF NOT EXISTS {table_name} ({columns})"
                )
                for row in data:
                    await db.execute(
                        f"INSERT OR REPLACE INTO {table_name} VALUES ({placeholders})",
                        list(row.values()),
                    )
                await db.commit()
        return str(db_path)

    async def _save_json(self, data: list[dict], path: Path) -> str:
        from json import dumps
        import aiofiles

        async with aiofiles.open(path, "w", encoding="utf-8") as f:
            await f.write(dumps(data, ensure_ascii=False, indent=2))
        return str(path)
